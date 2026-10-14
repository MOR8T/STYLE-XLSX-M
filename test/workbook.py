"""Independent synthetic fixture and workbook inspection (requires openpyxl)."""
import datetime
import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side


def generate(path):
    wb = openpyxl.Workbook()
    wb.properties.creator = 'Synthetic fixture'
    wb.properties.lastModifiedBy = 'Synthetic fixture'
    ws = wb.active
    ws.title = 'Template'
    ws['A1'] = 'Product A'
    ws['A1'].fill = PatternFill('solid', fgColor='FF12AB34')
    ws['B1'] = 'Product B'
    ws['B1'].fill = PatternFill('solid', fgColor=Color(theme=4, tint=0.4))
    ws['C1'] = 'Product C'
    ws['C1'].fill = PatternFill('solid', fgColor=Color(theme=0))
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, italic=True,
                         underline='single', color='FFAB1234')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='top',
                                  wrap_text=True, text_rotation=30)
    ws['A1'].border = Border(
        left=Side(style='thin', color='FF123456'),
        right=Side(style='thick', color='FF654321'),
        top=Side(style='dashed', color=Color(theme=4, tint=0.2)),
        bottom=Side(style='double', color=Color(indexed=10)))
    ws['A2'] = 100.25
    ws['A2'].number_format = '#,##0.00'
    ws['B2'] = datetime.datetime(2026, 1, 1)
    ws['B2'].number_format = 'yyyy-mm-dd'
    ws.merge_cells('A4:C4')
    ws['A4'] = 'Synthetic heading'
    ws.column_dimensions['A'].width = 24
    ws.row_dimensions[1].height = 32
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.scale = 91
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def color(c):
    if c is None:
        return None
    # RGB is compared without alpha: the writer normalizes six digit RGB to ARGB.
    value = c.value[-6:] if c.type == 'rgb' else c.value
    return {'type': c.type, 'value': value, 'tint': c.tint}


def inspect(path):
    ws = openpyxl.load_workbook(path).active
    a = ws['A1']
    result = {}
    for address in ('A1', 'B1', 'C1'):
        result[address + '.fill.patternType'] = ws[address].fill.patternType
        result[address + '.fill.fgColor'] = color(ws[address].fill.fgColor)
    for side in ('left', 'right', 'top', 'bottom'):
        b = getattr(a.border, side)
        result['A1.border.' + side + '.style'] = b.style
        result['A1.border.' + side + '.color'] = color(b.color)
    for prop in ('name', 'sz', 'bold', 'italic', 'underline'):
        result['A1.font.' + prop] = getattr(a.font, prop)
    result['A1.font.color'] = color(a.font.color)
    for prop in ('horizontal', 'vertical', 'wrapText', 'textRotation'):
        result['A1.alignment.' + prop] = getattr(a.alignment, prop)
    for address in ('A2', 'B2'):
        result[address + '.number_format'] = ws[address].number_format
    result['merges'] = sorted(str(r) for r in ws.merged_cells.ranges)
    result['column.A.width'] = ws.column_dimensions['A'].width
    result['row.1.height'] = ws.row_dimensions[1].height
    for prop in ('paperSize', 'orientation', 'scale'):
        result['pageSetup.' + prop] = getattr(ws.page_setup, prop)
    result['A1.value'] = a.value
    return result


def check_examples(directory):
    d = Path(directory)
    ws = openpyxl.load_workbook(d / 'demo.xlsx').active
    assert ws['A1'].alignment.wrapText is True
    assert ws['B1'].font.name == 'Courier' and ws['B1'].font.sz == 24
    assert ws['C1'].fill.patternType == 'solid'
    assert color(ws['C1'].fill.fgColor)['value'] == 'E9E9E9'
    assert ws['D1'].font.bold and color(ws['D1'].font.color)['value'] == 'FF0000'
    edited = openpyxl.load_workbook(d / 'demo-edited.xlsx').active
    assert edited['A1'].value == 'changed' and edited['A1'].alignment.wrapText
    assert color(edited['C1'].fill.fgColor)['value'] == 'E9E9E9'
    ws = openpyxl.load_workbook(d / 'reference.xlsx').active
    c = ws['A1']
    assert c.alignment.horizontal == c.alignment.vertical == 'center'
    assert c.alignment.wrapText and c.alignment.textRotation == 0
    for side in ('top', 'bottom', 'left', 'right'):
        assert getattr(c.border, side).style == 'thick'
        assert color(getattr(c.border, side).color)['value'] == 'FF0000'
    assert c.border.diagonal.style == 'thin'
    assert c.border.diagonalDown and not c.border.diagonalUp
    assert c.font.name == 'Calibri' and c.font.sz == 11
    assert color(c.font.color)['value'] == 'FF0000'
    assert c.font.vertAlign == 'superscript'
    assert not any((c.font.bold, c.font.italic, c.font.strike, c.font.underline))
    assert c.fill.patternType == 'solid' and color(c.fill.fgColor)['value'] == 'FF0000'
    ws = openpyxl.load_workbook(d / 'page-setup.xlsx').active
    assert ws.page_setup.paperSize == '9'
    assert ws.page_setup.scale == 91
    assert ws.page_setup.fitToWidth == 0
    assert ws.page_setup.orientation == 'landscape'
    return {'examples': 'passed'}


if __name__ == '__main__':
    action, path = sys.argv[1:3]
    if action == 'generate':
        generate(path)
    else:
        print(json.dumps(inspect(path) if action == 'inspect' else check_examples(path)))
